
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


def extract_text(uploaded_file):
    """Extract text from a Streamlit-uploaded TXT, PDF, or DOCX file."""
    if uploaded_file is None:
        return ""

    filename = uploaded_file.name.lower()
    raw_bytes = uploaded_file.getvalue()

    if filename.endswith(".txt"):
        for encoding in ("utf-8", "utf-8-sig", "latin-1"):
            try:
                return raw_bytes.decode(encoding)
            except UnicodeDecodeError:
                continue
        return raw_bytes.decode("utf-8", errors="replace")

    if filename.endswith(".pdf"):
        reader = PdfReader(BytesIO(raw_bytes))
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        return "\n\n".join(pages)

    if filename.endswith(".docx"):
        document = Document(BytesIO(raw_bytes))
        parts = [paragraph.text for paragraph in document.paragraphs]

        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))

        return "\n".join(parts)

    raise ValueError(
        f"Unsupported file type: {Path(filename).suffix}. "
        "Upload a TXT, PDF, or DOCX file."
    )
